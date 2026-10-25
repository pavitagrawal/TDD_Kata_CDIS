from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Car_Dealership_Inventory_System_Project_Report.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead:
        set_font(p.add_run(bold_lead), bold=True)
    set_font(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_font(r, size=15 if level == 1 else 12.5, bold=True)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        set_font(p.add_run(item))


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(value), size=9.5, bold=True, color="FFFFFF")
        if widths:
            cell.width = Inches(widths[i])
    for row_num, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            if row_num % 2:
                set_cell_shading(cell, "EDF3F8")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(value), size=9.5)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_page_number(footer):
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Page ")
    set_font(r, size=9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)
    add_page_number(sec.footer)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(80)
    title.paragraph_format.space_after = Pt(15)
    set_font(title.add_run("Car Dealership Inventory System"), size=24, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(45)
    set_font(subtitle.add_run("Project Report"), size=15, color="1F4E78")
    for label, value in [
        ("Submitted by", "____________________________"),
        ("Roll Number", "____________________________"),
        ("Programme", "____________________________"),
        ("Institution", "____________________________"),
        ("Submission Date", "____________________________"),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        set_font(p.add_run(f"{label}: "), bold=True)
        set_font(p.add_run(value))
    doc.add_page_break()

    add_heading(doc, "Abstract")
    add_body(doc, "The Car Dealership Inventory System is a full-stack web application designed to manage vehicle inventory while providing a simple browsing and purchasing experience for authenticated users. The project combines a React and TypeScript client with a FastAPI backend, PostgreSQL persistence, Alembic database migrations, and token-based authentication. Its core business rule is safe stock handling: a purchase is accepted only when inventory is available, and the database update is designed to prevent two competing requests from consuming the final unit.")
    add_body(doc, "The solution applies a modular-monolith design. HTTP routes handle transport concerns, service modules contain business rules, schemas define validated API contracts, and SQLAlchemy models manage persistence. The frontend provides responsive inventory browsing, filters, stock-aware actions, authentication views, and an accessible vehicle-details dialog. The project is structured for a managed PostgreSQL deployment with a separately deployed static frontend and backend service.")

    add_heading(doc, "Keywords")
    add_body(doc, "FastAPI, React, TypeScript, PostgreSQL, SQLAlchemy, JWT, Argon2id, Alembic, Test Driven Development, inventory management.")

    add_heading(doc, "1 Introduction")
    add_body(doc, "Car dealerships require accurate inventory information for staff and prospective buyers. Manual lists and loosely connected tools make it difficult to identify available vehicles, apply consistent validation, and prevent stock from becoming negative when multiple requests occur at the same time. This project addresses those needs with a browser-based inventory system and a backend that enforces business rules centrally.")

    add_heading(doc, "2 Objectives")
    add_bullets(doc, [
        "Provide secure registration, login, and current-user identification.",
        "Allow authenticated users to browse, search, filter, and purchase available vehicles.",
        "Restrict vehicle creation, updates, deletion, and restocking to administrators.",
        "Validate input at the API boundary and preserve domain constraints in PostgreSQL.",
        "Use database migrations and an environment-based configuration model suitable for deployment.",
        "Demonstrate incremental testing and quality checks for the frontend and backend.",
    ])

    add_heading(doc, "3 Scope")
    add_body(doc, "The delivered scope covers user authentication, vehicle inventory, filtering, stock-aware purchasing, administrator-protected backend operations, a responsive user interface, and a deployment runbook. The project is not an e-commerce payment system. It does not process real payments, manage financing, or integrate with third-party dealer-management systems.")

    add_heading(doc, "4 Technology Stack")
    add_table(doc, ["Layer", "Technology", "Purpose"], [
        ("Frontend", "React 19, TypeScript, Vite", "Responsive single-page application and typed UI."),
        ("Backend", "Python 3.12+, FastAPI, Pydantic", "REST API, validation, and dependency injection."),
        ("Data", "PostgreSQL, SQLAlchemy 2, Alembic", "Relational storage, ORM mapping, and migrations."),
        ("Security", "PyJWT, pwdlib with Argon2id", "Expiring bearer tokens and password hashing."),
        ("Quality", "Vitest, Testing Library, Ruff, mypy, ESLint, Prettier", "Tests, static analysis, and formatting."),
        ("Delivery", "GitHub Actions, managed PostgreSQL, static host, Python web service", "Automated checks and planned deployment."),
    ], [1.0, 2.1, 3.1])

    add_heading(doc, "5 System Architecture")
    add_body(doc, "The application follows a modular-monolith architecture. The React client communicates with the FastAPI API over HTTPS. The API routes call service modules, which apply authorization and inventory rules before using SQLAlchemy to access PostgreSQL. Alembic tracks structural database changes. This separation keeps the initial solution simple while isolating responsibilities for future extension.")
    add_table(doc, ["Component", "Responsibility"], [
        ("React client", "Presents authentication, inventory filters, vehicle cards, purchase actions, and accessible details."),
        ("API routes", "Maps HTTP requests and responses; injects the current user and database session."),
        ("Services", "Implements authentication decisions, role checks, inventory search, purchase, and restocking logic."),
        ("Schemas", "Validates request and response data, including non-blank vehicle fields and valid prices."),
        ("PostgreSQL", "Stores users and vehicles, applies constraints, indexes searchable fields, and supports concurrency-safe updates."),
    ], [1.45, 4.75])

    add_heading(doc, "6 Functional Requirements")
    add_table(doc, ["ID", "Requirement", "Implementation Status"], [
        ("FR1", "Users can register and log in.", "Implemented"),
        ("FR2", "Authenticated users can view and filter inventory.", "Implemented"),
        ("FR3", "Authenticated users can purchase an available vehicle.", "Implemented"),
        ("FR4", "Administrators can create, edit, delete, and restock vehicles through the API.", "Implemented"),
        ("FR5", "The interface presents vehicle details and disables purchase for unavailable stock.", "Implemented"),
        ("FR6", "A dedicated administrator inventory workspace in the frontend.", "Planned improvement"),
    ], [0.55, 4.3, 1.35])

    add_heading(doc, "7 Data Design and Business Rules")
    add_body(doc, "The system stores users and vehicles. A user has a unique normalized email address, password hash, role, and timestamps. A vehicle has make, model, category, price, stock quantity, and timestamps. Database constraints require a non-negative price and stock quantity. Request schemas trim vehicle text fields and reject values that are blank after trimming.")
    add_body(doc, "The purchase operation uses a conditional database update equivalent to updating a vehicle only when its stock quantity is greater than zero. This moves the critical stock check into a single statement, reducing the risk of an oversell caused by simultaneous requests. The complete concurrent behavior must be verified against a running PostgreSQL instance before production release.")

    add_heading(doc, "8 Security Design")
    add_bullets(doc, [
        "Passwords are hashed with Argon2id through pwdlib; plaintext passwords are not stored.",
        "JWT bearer tokens include an expiry and are signed using an environment-provided secret of at least 32 characters.",
        "The client keeps the bearer token in memory, reducing persistence across browser refreshes.",
        "Authorization is enforced by backend roles, not by hiding client-side controls.",
        "SQLAlchemy parameterizes database expressions, avoiding string concatenation of user input into SQL.",
        "CORS uses an explicit environment-configured allow-list for the deployed frontend origin.",
    ])
    add_body(doc, "The present security posture should be extended with authentication rate limiting, production monitoring, and penetration testing before a public production release.")

    add_heading(doc, "9 API Summary")
    add_table(doc, ["Endpoint", "Access", "Purpose"], [
        ("POST /api/auth/register", "Public", "Create a standard user account."),
        ("POST /api/auth/login", "Public", "Return an expiring bearer token."),
        ("GET /api/auth/me", "Authenticated", "Return the current user."),
        ("GET /api/vehicles and /search", "Authenticated", "Browse and filter inventory."),
        ("POST, PUT, DELETE /api/vehicles", "Administrator", "Manage vehicle inventory."),
        ("POST /api/vehicles/{id}/purchase", "Authenticated", "Purchase one available unit."),
        ("POST /api/vehicles/{id}/restock", "Administrator", "Increase stock quantity."),
    ], [2.55, 1.15, 2.5])

    add_heading(doc, "10 Testing and Quality Assurance")
    add_body(doc, "The project was developed with an incremental test-first intent. The frontend contains Vitest and Testing Library checks for core rendered behavior, including the stock-disabled purchase state. The backend includes tests for authentication utilities, schema validation, configuration validation, and vehicle service behavior. The continuous-integration workflow is configured to run frontend checks and PostgreSQL-backed backend checks when dependencies and services are available.")
    add_table(doc, ["Check", "Observed Local Result"], [
        ("Frontend lint", "Passed"),
        ("Frontend formatting", "Passed"),
        ("Frontend unit tests", "Passed: 2 test files and 2 tests"),
        ("Frontend production build", "Passed"),
        ("Python syntax compilation", "Passed"),
        ("Backend tests, migrations, coverage, type checks", "Not executed locally because PyPI dependency installation timed out and PostgreSQL/Docker were unavailable"),
        ("npm audit", "Registry audit endpoint was unavailable in the local environment"),
    ], [2.5, 3.7])

    add_heading(doc, "11 Deployment Plan")
    add_body(doc, "The planned release uses a managed PostgreSQL database, a Python web-service host for the API, and a static host for the Vite frontend. The production backend requires DATABASE_URL, JWT_SECRET, APP_ENV=production, and a JSON CORS allow-list. The frontend requires VITE_API_BASE_URL at build time. Database migration is run once against the production database before the application is released.")
    add_bullets(doc, [
        "Create the managed PostgreSQL database and retain its connection string as a secret.",
        "Deploy the FastAPI backend and configure a GET /health health check.",
        "Run alembic upgrade head against the production database.",
        "Deploy the frontend with the public API origin configured at build time.",
        "Set the backend CORS allow-list to the final HTTPS frontend origin.",
        "Smoke-test registration, login, search, purchase, and an administrator action over HTTPS.",
    ])

    add_heading(doc, "12 Limitations and Future Enhancements")
    add_bullets(doc, [
        "Complete the administrator-facing create, edit, delete, and restock workspace in the frontend.",
        "Run and record PostgreSQL integration tests, migrations, rollback tests, and concurrency tests.",
        "Add authentication rate limiting, audit logging, monitoring, and error tracking.",
        "Add richer inventory media, pagination, saved searches, and administrative reporting.",
        "Capture final screenshots and CI evidence from the deployed application.",
    ])

    add_heading(doc, "13 Conclusion")
    add_body(doc, "The Car Dealership Inventory System establishes a solid full-stack foundation for inventory management. It separates interface, API, business-rule, and persistence responsibilities; validates data at multiple layers; and keeps authorization on the server. The frontend has passed its local lint, formatting, test, and build checks. Before deployment, the remaining backend verification and PostgreSQL integration checks should be executed in an environment with network access and a running database. Completing the administrator workspace and production verification will make the project ready for a stronger demonstration and submission.")

    add_heading(doc, "References")
    add_bullets(doc, [
        "FastAPI documentation. https://fastapi.tiangolo.com/",
        "React documentation. https://react.dev/",
        "PostgreSQL documentation. https://www.postgresql.org/docs/",
        "SQLAlchemy documentation. https://docs.sqlalchemy.org/",
        "Alembic documentation. https://alembic.sqlalchemy.org/",
    ])

    doc.core_properties.title = "Car Dealership Inventory System Project Report"
    doc.core_properties.subject = "Software project report"
    doc.core_properties.author = ""
    doc.core_properties.comments = ""
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
